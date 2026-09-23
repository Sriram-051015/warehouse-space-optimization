import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
import os

# ============================================================
# FILE PATHS
# ============================================================

analytics_file = "data/processed/final_warehouse_analytics.csv"
space_file = "results/space_utilization_summary.csv"

output_file = (
    "results/report/"
    "warehouse_space_optimization_report.docx"
)

print("\n===== CREATING OPTIMIZATION REPORT =====")

# ============================================================
# LOAD DATA
# ============================================================

data = pd.read_csv(analytics_file)
space = pd.read_csv(space_file)

# ============================================================
# CALCULATE PROJECT METRICS
# ============================================================

total_products = len(data)

total_volume = data["product_volume"].sum()

high_priority = (
    data["storage_priority"] == "High Priority"
).sum()

medium_priority = (
    data["storage_priority"] == "Medium Priority"
).sum()

low_priority = (
    data["storage_priority"] == "Low Priority"
).sum()

relocation_count = (
    data["relocation_recommended"] == True
).sum()

# ============================================================
# CREATE DOCUMENT
# ============================================================

document = Document()

# ------------------------------------------------------------
# TITLE
# ------------------------------------------------------------

title = document.add_heading(
    "Warehouse Space Optimization",
    level=0
)

title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = document.add_paragraph(
    "AI and Data Analytics Based Product Placement Optimization"
)

subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ============================================================
# 1. PROJECT OBJECTIVE
# ============================================================

document.add_heading(
    "1. Project Objective",
    level=1
)

document.add_paragraph(
    "The objective of this project is to analyze warehouse "
    "product characteristics, movement frequency, and existing "
    "warehouse storage locations to support data-driven product "
    "placement recommendations and improve warehouse space "
    "planning."
)

# ============================================================
# 2. INPUT DATA
# ============================================================

document.add_heading(
    "2. Input Data",
    level=1
)

document.add_paragraph(
    "The project uses an external warehouse dataset containing "
    "product information and historical warehouse pick activity."
)

inputs = [
    "Product dimensions: height, width, and depth.",
    "Movement frequency: historical product pick frequency.",
    "Warehouse layout information: existing product storage "
    "locations recorded in the pick history."
]

for item in inputs:
    document.add_paragraph(
        item,
        style="List Bullet"
    )

# ============================================================
# 3. DATA PREPROCESSING
# ============================================================

document.add_heading(
    "3. Data Preprocessing",
    level=1
)

document.add_paragraph(
    "The raw SKU and pick-history files were inspected and "
    "processed using Python and pandas. The SKU file contained "
    "product dimensions, while the pick-history file was "
    "processed in chunks to handle its large size."
)

document.add_paragraph(
    "Product volume was calculated from the first-unit dimensions "
    "using height, width, and depth. Pick records were aggregated "
    "to calculate product movement frequency."
)

# ============================================================
# 4. MOVEMENT ANALYSIS
# ============================================================

document.add_heading(
    "4. Movement Frequency Analysis",
    level=1
)

document.add_paragraph(
    f"The final analytical dataset contains {total_products:,} "
    "products with usable dimension and movement information."
)

document.add_paragraph(
    "Products were categorized into Low Movement, Medium Movement, "
    "and High Movement groups using percentile-based thresholds."
)

movement_counts = (
    data["movement_category"]
    .value_counts()
)

for category in [
    "Low Movement",
    "Medium Movement",
    "High Movement"
]:
    count = movement_counts.get(category, 0)

    document.add_paragraph(
        f"{category}: {count:,} products",
        style="List Bullet"
    )

# ============================================================
# 5. CLUSTERING
# ============================================================

document.add_heading(
    "5. Product Clustering",
    level=1
)

document.add_paragraph(
    "K-Means clustering was applied using product volume and "
    "movement frequency as the main clustering features. "
    "Four clusters were created to identify groups of products "
    "with similar space and movement characteristics."
)

document.add_paragraph(
    "StandardScaler was used before clustering so that product "
    "volume and movement frequency could contribute to the "
    "clustering process on comparable scales."
)

# ============================================================
# 6. STORAGE PRIORITY
# ============================================================

document.add_heading(
    "6. Storage Priority Calculation",
    level=1
)

document.add_paragraph(
    "A storage priority score was calculated using a weighted "
    "combination of movement frequency and product space "
    "requirement."
)

document.add_paragraph(
    "Movement frequency received a 70% weight and product "
    "space requirement received a 30% weight. Products were "
    "then classified into High Priority, Medium Priority, "
    "and Low Priority groups."
)

document.add_paragraph(
    f"High Priority: {high_priority:,} products",
    style="List Bullet"
)

document.add_paragraph(
    f"Medium Priority: {medium_priority:,} products",
    style="List Bullet"
)

document.add_paragraph(
    f"Low Priority: {low_priority:,} products",
    style="List Bullet"
)

# ============================================================
# 7. WAREHOUSE LOCATION ANALYSIS
# ============================================================

document.add_heading(
    "7. Warehouse Location Analysis",
    level=1
)

document.add_paragraph(
    "Historical FROM_LOC values from the warehouse pick history "
    "were analyzed to identify the main recorded storage location "
    "for each SKU."
)

document.add_paragraph(
    "The most frequently used recorded location for each product "
    "was retained as its main storage location."
)

# ============================================================
# 8. PLACEMENT RECOMMENDATION
# ============================================================

document.add_heading(
    "8. Placement Recommendation",
    level=1
)

document.add_paragraph(
    "Products were assigned to analytical storage zones based "
    "on their calculated storage priority."
)

document.add_paragraph(
    "ZONE_A represents High Priority products, ZONE_B represents "
    "Medium Priority products, and ZONE_C represents Low Priority "
    "products."
)

document.add_paragraph(
    f"The analysis identifies {relocation_count:,} products as "
    "requiring a relocation recommendation under this analytical "
    "zoning approach."
)

# ============================================================
# 9. SPACE ANALYSIS
# ============================================================

document.add_heading(
    "9. Space Utilization Analysis",
    level=1
)

document.add_paragraph(
    f"The total calculated product volume in the final dataset "
    f"is {total_volume:,.0f} cubic units."
)

for _, row in space.iterrows():

    document.add_paragraph(
        f"{row['recommended_zone']}: "
        f"{int(row['product_count']):,} products, "
        f"{row['space_percentage']:.2f}% of total product volume.",
        style="List Bullet"
    )

document.add_paragraph(
    "The space percentages represent the distribution of "
    "calculated product volume across the recommended zones. "
    "They should not be interpreted as physical warehouse "
    "capacity utilization because the source dataset does not "
    "provide complete bin-capacity or warehouse-capacity data."
)

# ============================================================
# 10. DASHBOARD
# ============================================================

document.add_heading(
    "10. Analytics Dashboard",
    level=1
)

document.add_paragraph(
    "A Power BI dashboard was created to visualize the analysis."
)

dashboard_items = [
    "Product volume versus movement frequency.",
    "Movement frequency by recommended zone.",
    "Product space by recommended zone.",
    "Product distribution by movement category.",
    "Frequently used warehouse storage locations.",
    "Recommended storage zones by priority.",
    "Key performance indicators for products, product volume, "
    "high-priority products, and relocation recommendations."
]

for item in dashboard_items:
    document.add_paragraph(
        item,
        style="List Bullet"
    )

# ============================================================
# 11. PROJECT LIMITATIONS
# ============================================================

document.add_heading(
    "11. Project Limitations",
    level=1
)

limitations = [
    "The source data does not provide complete physical bin "
    "capacity for every warehouse location.",
    "The recommended zones are analytical zones and are not "
    "literal physical warehouse coordinates.",
    "Exact aisle distance and travel-time information was not "
    "available in the processed dataset.",
    "The placement recommendation should therefore be treated "
    "as a decision-support analysis rather than an automated "
    "physical warehouse assignment system."
]

for item in limitations:
    document.add_paragraph(
        item,
        style="List Bullet"
    )

# ============================================================
# 12. CONCLUSION
# ============================================================

document.add_heading(
    "12. Conclusion",
    level=1
)

document.add_paragraph(
    "This project demonstrates how warehouse product dimensions, "
    "historical movement frequency, and existing storage-location "
    "data can be combined with clustering and data analytics to "
    "support warehouse space optimization."
)

document.add_paragraph(
    f"The final analysis covers {total_products:,} products and "
    f"calculates a total product volume of {total_volume:,.0f} "
    "cubic units. The resulting priority classification and "
    "analytical zone recommendations provide a structured basis "
    "for warehouse placement decisions."
)

# ============================================================
# 13. TECHNOLOGIES USED
# ============================================================

document.add_heading(
    "13. Technologies Used",
    level=1
)

technologies = [
    "Python",
    "Pandas",
    "NumPy",
    "Scikit-learn",
    "K-Means Clustering",
    "Matplotlib",
    "Seaborn",
    "Power BI",
    "Git and GitHub"
]

for technology in technologies:
    document.add_paragraph(
        technology,
        style="List Bullet"
    )

# ============================================================
# SAVE DOCUMENT
# ============================================================

os.makedirs(
    "results/report",
    exist_ok=True
)

document.save(output_file)

print("\n===== REPORT CREATED =====")

print(output_file)

print("\n===== REPORT SUMMARY =====")

print("Total products:", total_products)
print("Total product volume:", total_volume)
print("High priority:", high_priority)
print("Medium priority:", medium_priority)
print("Low priority:", low_priority)
print("Relocation recommendations:", relocation_count)

print("\n===== OPTIMIZATION REPORT COMPLETE =====")